import os
import requests
import json
from docx import Document
from docx.shared import Inches
import base64
from PIL import Image
from io import BytesIO, StringIO

def ask_llm(encoded_image):
    url = f'http://localhost:11435/api/generate'
    data = {
        "model": "minicpm-v",
        "prompt": "Describe this sequence diagram, found in a 3GPP standard. Elicit as much specific information as possible.",
        "images": [encoded_image],
        "stream": False
        }
    headers = {'Content-Type': 'application/json'}
    response = requests.post(url, data=json.dumps(data), headers=headers)
    json_data = json.loads(response.text)
    return json_data['response']

def extract_images_from_docx(docx_path, output_folder):
    input_doc = Document(docx_path)
    output_doc = Document()
    current_section = None
    rels = input_doc.part.rels

    for paragraphs in input_doc.paragraphs:
        if paragraphs.style.name.startswith('Heading'):
            current_section = paragraphs.text
            print(current_section)
        
        for rel in rels:
            if "image" in rels[rel].target_ref:
                image = rels[rel].target_part.blob
                img_name = os.path.basename(rels[rel].target_ref)
                img_path = os.path.join(output_folder, img_name)
                # print(img_path)
                # with open(img_path, "wb") as img_file:
                #     img_file.write(image)
                # encode_images_to_base64(output_folder)
                if current_section:
                        output_doc.add_heading(current_section, level=1)
                        current_section = None # Reset current_section
 
                # take the image and add it to the docx
                # img_path = os.path.join(output_folder, img_name)
                # output_doc.add_picture(img_path, width=Inches(6))
                # print(f"Added image {img_name} to the document.")

                # encoded_image = base64.b64encode(image).decode('utf-8') 
                bytes = BytesIO(image)    
                img = Image.open(StringIO(bytes))  
                output_doc.add_picture(img, width=Inches(6))
                print(type(img))

        output_doc.save(output_doc)
                    



def encode_images_to_base64(output_folder):
    for img_name in os.listdir(output_folder):
        img_path = os.path.join(output_folder, img_name)
        if img_name.endswith(".emf"):
            new_name = img_name.replace(".emf", ".png")
            new_path = os.path.join(output_folder, new_name)
            Image.open(img_path).save(output_folder + "/" + new_name)
            with open(new_path, "rb") as img_file:
                encoded_string = base64.b64encode(img_file.read()).decode('utf-8')
                print(f"Finished encoding {new_name}.")

                llm_response = ask_llm(encoded_string)

                txt_path = os.path.join(output_folder, f"{new_name[0:-4]}.txt")
                with open(txt_path, "w") as txt_file:
                    txt_file.write(llm_response)

def main():
    standards_folder = "test_files"
    output_folder = "extracted_images"
    output_file = "extracted_images.docx"
    
    # if not os.path.exists(output_folder):
    #     os.makedirs(output_folder)
    
    # for filename in os.listdir(standards_folder):
    #     if filename.endswith(".docx"):
    #         docx_path = os.path.join(standards_folder, filename)
    #         extract_images_from_docx(docx_path, output_folder)
    #         print(f"Extracted images from {filename}")

    # encode_images_to_base64(output_folder)

    print("Processing images")
    extract_images_from_docx("test_files/23502-i20_l.docx", output_file)
    

if __name__ == "__main__":
    main()